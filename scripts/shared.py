from pathlib import Path
from typing import List, Tuple


SOURCE_DIR = Path(__file__).resolve().parent.parent / "source"


def get_all_files() -> List[Path]:
    return sorted(SOURCE_DIR.rglob("*"))


def get_all_files_with_ext(extensions: set) -> List[Path]:
    return sorted(
        p for p in SOURCE_DIR.rglob("*")
        if p.is_file() and p.suffix.lower() in extensions
    )


def extract_text(filepath: Path) -> str:
    suffix = filepath.suffix.lower()
    try:
        if suffix == ".docx":
            return _extract_docx(filepath)
        elif suffix == ".pdf":
            return _extract_pdf(filepath)
        elif suffix == ".xlsx":
            return _extract_xlsx(filepath)
        elif suffix == ".doc":
            return _extract_doc(filepath)
        elif suffix in (".png", ".jpg", ".jpeg", ".gif", ".bmp"):
            return ""
        else:
            return ""
    except Exception as e:
        print(f"  [WARN] Error extracting text from {filepath.name}: {e}")
        return ""


def _extract_docx(filepath: Path) -> str:
    from docx import Document
    doc = Document(str(filepath))
    paragraphs = [p.text for p in doc.paragraphs]
    return "\n".join(paragraphs)


def _extract_pdf(filepath: Path) -> str:
    import pdfplumber
    text_parts = []
    with pdfplumber.open(str(filepath)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def _extract_xlsx(filepath: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(filepath), read_only=True, data_only=True)
    text_parts = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = " ".join(str(cell) for cell in row if cell is not None)
            if row_text.strip():
                text_parts.append(row_text)
    return "\n".join(text_parts)


def _extract_doc(filepath: Path) -> str:
    raise RuntimeError(
        f".doc format not supported (require catdoc/antiword): {filepath.name}"
    )
